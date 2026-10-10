import csv
import io
import os
from datetime import timedelta

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import SetPasswordForm
from django.core.exceptions import PermissionDenied
from django.core.mail import EmailMessage
from django.db.models import Count, Prefetch, Q
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from . import importer
from .forms import (
    AttachmentForm,
    CalibrationEventForm,
    DeviceFilterForm,
    DeviceForm,
    DeviceModelEditForm,
    DeviceModelForm,
    FieldDefinitionForm,
    ImportUploadForm,
    ManufacturerEditForm,
    ManufacturerForm,
    ProfileForm,
    SiteSettingsForm,
    UserCreateForm,
    UserEditForm,
)
from .ics import build_calendar
from .models import (
    Attachment,
    AuditLog,
    CalibrationEvent,
    Device,
    DeviceFieldValue,
    DeviceModel,
    FieldDefinition,
    Manufacturer,
    SiteSettings,
    User,
    new_ics_token,
)
from .permissions import (
    admin_required,
    can_delete_device,
    can_edit,
    can_view,
    editor_all_required,
    visible_devices,
)
from .services import (
    audit,
    delete_calibration_event,
    form_changes,
    save_calibration_event,
    snapshot,
)


def _get_device(request, pk, edit=False):
    device = get_object_or_404(
        Device.objects.select_related("owner", "manufacturer", "device_model"), pk=pk
    )
    allowed = can_edit(request.user, device) if edit else can_view(request.user, device)
    if not allowed:
        raise PermissionDenied
    return device


def _file_response(fieldfile, filename):
    try:
        return FileResponse(fieldfile.open("rb"), as_attachment=True, filename=filename)
    except FileNotFoundError:
        raise Http404("File is missing on the server.")


# ---------------------------------------------------------------- dashboard


@login_required
def dashboard(request):
    today = timezone.localdate()
    settings = SiteSettings.load()
    soon = today + timedelta(days=settings.due_soon_days)
    qs = visible_devices(request.user).select_related("owner", "manufacturer", "device_model")
    active = qs.filter(status__in=Device.REMINDER_STATUSES)
    context = {
        "today": today,
        "settings": settings,
        "overdue": active.filter(next_calibration_date__lt=today).order_by("next_calibration_date"),
        "due_soon": active.filter(next_calibration_date__gte=today, next_calibration_date__lte=soon)
        .order_by("next_calibration_date"),
        "never": active.filter(next_calibration_date__isnull=True),
        "at_calibration": qs.filter(status=Device.STATUS_AT_CALIBRATION),
        "ok_count": active.filter(next_calibration_date__gt=soon).count(),
        "total": qs.exclude(status=Device.STATUS_RETIRED).count(),
        "my_count": Device.objects.filter(owner=request.user).exclude(status=Device.STATUS_RETIRED).count(),
    }
    if request.user.can_view_all:
        context["recent"] = AuditLog.objects.select_related("device")[:10]
    return render(request, "calibration/dashboard.html", context)


# ------------------------------------------------------------------ devices


def _filtered_devices(request):
    form = DeviceFilterForm(request.GET or None, user=request.user)
    qs = visible_devices(request.user).select_related("owner", "manufacturer", "device_model")
    if form.is_valid():
        d = form.cleaned_data
        if d["q"]:
            q = d["q"]
            qs = qs.filter(
                Q(asset_tag__icontains=q) | Q(name__icontains=q) | Q(serial_number__icontains=q)
                | Q(location__icontains=q) | Q(manufacturer__name__icontains=q) | Q(device_model__name__icontains=q)
                | Q(field_values__value_text__icontains=q)
            ).distinct()
        if d["status"] == "in_use":
            qs = qs.exclude(status=Device.STATUS_RETIRED)
        elif d["status"]:
            qs = qs.filter(status=d["status"])
        if d.get("owner"):
            qs = qs.filter(owner=d["owner"])
        if d["manufacturer"]:
            qs = qs.filter(manufacturer=d["manufacturer"])
        if d["device_model"]:
            qs = qs.filter(device_model=d["device_model"])
        if d["due"]:
            today = timezone.localdate()
            soon = today + timedelta(days=SiteSettings.load().due_soon_days)
            active = Q(status__in=Device.REMINDER_STATUSES)
            qs = qs.filter({
                "overdue": active & Q(next_calibration_date__lt=today),
                "due_soon": active & Q(next_calibration_date__gte=today, next_calibration_date__lte=soon),
                "ok": active & Q(next_calibration_date__gt=soon),
                "unknown": active & Q(next_calibration_date__isnull=True),
            }[d["due"]])
    return form, qs


def _list_columns(form):
    """show_in_list fields of the manufacturer and model being filtered on."""
    if not form.is_valid():
        return []
    brand, model = form.cleaned_data["manufacturer"], form.cleaned_data["device_model"]
    if model and not brand:
        brand = model.manufacturer
    owners = [o for o in (brand, model) if o]
    return [f for o in owners for f in o.fields.filter(show_in_list=True)]


@login_required
def device_list(request):
    form, qs = _filtered_devices(request)
    fmt = request.GET.get("export")
    if fmt in ("csv", "xlsx"):
        return export_devices(qs, fmt)
    extra_fields = _list_columns(form)
    devices = list(qs.prefetch_related("field_values"))
    soon_days = SiteSettings.load().due_soon_days
    today = timezone.localdate()
    rows = []
    for d in devices:
        values = {v.field_id: v for v in d.field_values.all()}
        rows.append({
            "device": d,
            "state": d.due_state(today, soon_days),
            "days": d.days_until_due(today),
            "extra": [values[f.id].display() if f.id in values else "" for f in extra_fields],
        })
    query = request.GET.copy()
    query.pop("export", None)
    return render(request, "calibration/device_list.html", {
        "form": form, "rows": rows, "extra_fields": extra_fields, "query": query.urlencode(),
    })


EXPORT_COLUMNS = [
    ("Asset ID", lambda d: d.asset_tag),
    ("Name", lambda d: d.name),
    ("Owner", lambda d: d.owner.username),
    ("Owner name", lambda d: str(d.owner)),
    ("Status", lambda d: d.get_status_display()),
    ("Manufacturer", lambda d: d.manufacturer.name if d.manufacturer else ""),
    ("Model", lambda d: d.device_model.name if d.device_model else ""),
    ("Serial number", lambda d: d.serial_number),
    ("Location", lambda d: d.location),
    ("Interval (months)", lambda d: d.effective_interval_months),
    ("Interval override (months)", lambda d: d.interval_override_months or ""),
    ("Last calibration date", lambda d: d.last_calibration_date),
    ("Next calibration date", lambda d: d.next_calibration_date),
    ("Notes", lambda d: d.notes),
]


def export_devices(qs, fmt):
    devices = list(qs.prefetch_related("field_values__field", "manufacturer__fields", "device_model__fields"))
    custom_names = []
    for d in devices:
        for f in d.field_definitions():
            if f.field_type != FieldDefinition.FILE and f.name not in custom_names:
                custom_names.append(f.name)
    header = [c[0] for c in EXPORT_COLUMNS] + custom_names
    rows = []
    for d in devices:
        current = {f.id for f in d.field_definitions()}
        values = {v.field.name: v for v in d.field_values.all() if v.field_id in current}
        row = [fn(d) for _, fn in EXPORT_COLUMNS]
        for name in custom_names:
            v = values.get(name)
            row.append(v.value if v is not None and v.value is not None else "")
        rows.append(row)
    stamp = timezone.localdate().strftime("%Y%m%d")
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(header)
        for row in rows:
            w.writerow([x.strftime("%d.%m.%Y") if hasattr(x, "strftime") else x for x in row])
        resp = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="devices_{stamp}.csv"'
        return resp
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Devices"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([float(x) if hasattr(x, "as_tuple") else x for x in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if hasattr(cell.value, "year"):
                cell.number_format = "DD.MM.YYYY"
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp["Content-Disposition"] = f'attachment; filename="devices_{stamp}.xlsx"'
    return resp


@login_required
def device_detail(request, pk):
    device = _get_device(request, pk)
    today = timezone.localdate()
    settings = SiteSettings.load()
    return render(request, "calibration/device_detail.html", {
        "device": device,
        "custom_values": device.custom_values(),
        "events": device.events.select_related("created_by").prefetch_related("attachments"),
        "attachments": device.attachments.filter(event__isnull=True),
        "audit": device.audit.all()[:50],
        "state": device.due_state(today, settings.due_soon_days),
        "days": device.days_until_due(today),
        "can_edit": can_edit(request.user, device),
        "can_delete": can_delete_device(request.user),
        "attachment_form": AttachmentForm(),
    })


def _lookups_from_request(request, device=None):
    """Manufacturer and model chosen on the (re)submitted form; their fields are shown."""
    data = request.POST if request.method == "POST" else request.GET
    if "manufacturer" not in data and "device_model" not in data:
        return (device.manufacturer, device.device_model) if device else (None, None)

    def pick(model, key):
        raw = data.get(key)
        return model.objects.filter(pk=raw).first() if raw and raw.isdigit() else None

    model = pick(DeviceModel, "device_model")
    brand = pick(Manufacturer, "manufacturer") or (model.manufacturer if model else None)
    return brand, model


DEVICE_FIELDS = DeviceForm.Meta.fields


@login_required
def device_create(request):
    brand, model = _lookups_from_request(request)
    lookups = {"manufacturer": brand, "device_model": model}
    if request.method == "POST" and "_reload" not in request.POST:
        form = DeviceForm(request.POST, request.FILES, user=request.user, **lookups)
        if form.is_valid():
            device = form.save(commit=False)
            if not request.user.can_edit_all:
                device.owner = request.user
            device.save()
            form.save_custom_values(device)
            audit(request.user, "created", device, changes={
                str(form.fields[n].label or n): (None, form.cleaned_data.get(n))
                for n in DEVICE_FIELDS if n in form.cleaned_data and form.cleaned_data.get(n) not in (None, "")
            })
            first = form.cleaned_data.get("initial_calibration_date")
            if first:
                event = CalibrationEvent(
                    device=device, date=first, result=CalibrationEvent.RESULT_NOT_RECORDED,
                    notes="Entered when the device was created", created_by=request.user,
                )
                save_calibration_event(event, request.user, created=True)
            messages.success(request, f"Device {device.asset_tag} created.")
            return redirect("device_detail", pk=device.pk)
    else:
        initial = request.POST.dict() if request.method == "POST" else {"owner": request.user.pk}
        initial.pop("csrfmiddlewaretoken", None)
        form = DeviceForm(user=request.user, initial=initial, **lookups)
    return render(request, "calibration/device_form.html", {"form": form, "device": None})


@login_required
def device_edit(request, pk):
    device = _get_device(request, pk, edit=True)
    brand, model = _lookups_from_request(request, device)
    lookups = {"manufacturer": brand, "device_model": model}
    if request.method == "POST" and "_reload" not in request.POST:
        before = snapshot(device, DEVICE_FIELDS)
        form = DeviceForm(request.POST, request.FILES, instance=device, user=request.user, **lookups)
        if form.is_valid():
            changes = form_changes(form, before)
            device = form.save()
            device.refresh_calibration_dates()  # interval may have changed
            changes.update(form.save_custom_values(device))
            if changes:
                audit(request.user, "edited", device, changes=changes)
            messages.success(request, "Device saved.")
            return redirect("device_detail", pk=device.pk)
    else:
        initial = {}
        if request.method == "POST":
            initial = request.POST.dict()
            initial.pop("csrfmiddlewaretoken", None)
        form = DeviceForm(instance=device, user=request.user, initial=initial, **lookups)
    return render(request, "calibration/device_form.html", {"form": form, "device": device})


@login_required
def device_delete(request, pk):
    if not can_delete_device(request.user):
        raise PermissionDenied
    device = get_object_or_404(Device, pk=pk)
    if request.method == "POST":
        audit(request.user, "deleted", device, device=None, changes={"Asset ID": (device.asset_tag, None)})
        for att in device.attachments.all():
            att.file.delete(save=False)
        for v in device.field_values.exclude(value_file=""):
            v.value_file.delete(save=False)
        tag = device.asset_tag
        device.delete()
        messages.success(request, f"Device {tag} deleted.")
        return redirect("device_list")
    return render(request, "calibration/confirm_delete.html", {
        "object": device, "what": "device",
        "warning": "This removes the device with its full calibration history and files. "
                   "Consider setting the status to Retired instead.",
        "cancel_url": reverse("device_detail", args=[device.pk]),
    })


# ------------------------------------------------------- calibration events


def _save_event_attachments(request, form, event):
    for f in form.cleaned_data.get("files") or []:
        Attachment.objects.create(
            device=event.device, event=event, file=f, original_name=os.path.basename(f.name),
            uploaded_by=request.user,
        )


@login_required
def event_create(request, pk):
    device = _get_device(request, pk, edit=True)
    event = CalibrationEvent(device=device, created_by=request.user)
    if request.method == "POST":
        form = CalibrationEventForm(request.POST, request.FILES, instance=event)
        if form.is_valid():
            event = form.save(commit=False)
            save_calibration_event(event, request.user, created=True)
            _save_event_attachments(request, form, event)
            if device.status == Device.STATUS_AT_CALIBRATION and request.POST.get("set_active"):
                device.status = Device.STATUS_ACTIVE
                device.save(update_fields=["status", "updated_at"])
                audit(request.user, "edited", device, changes={"Status": ("At calibration", "Active")})
            messages.success(
                request,
                f"Calibration recorded. Next calibration: {device.next_calibration_date:%d.%m.%Y}."
                if device.next_calibration_date else "Calibration recorded.",
            )
            return redirect("device_detail", pk=device.pk)
    else:
        form = CalibrationEventForm(instance=event, initial={"date": timezone.localdate()})
    return render(request, "calibration/event_form.html", {"form": form, "device": device, "event": None})


@login_required
def event_edit(request, pk):
    event = get_object_or_404(CalibrationEvent.objects.select_related("device"), pk=pk)
    device = _get_device(request, event.device_id, edit=True)
    if request.method == "POST":
        before = snapshot(event, CalibrationEventForm.Meta.fields)
        form = CalibrationEventForm(request.POST, request.FILES, instance=event)
        if form.is_valid():
            changes = form_changes(form, before)
            changes.pop(str(form.fields["files"].label), None)
            event = form.save()
            device.refresh_calibration_dates()
            _save_event_attachments(request, form, event)
            if changes:
                audit(request.user, "calibration edited", event, device=device, changes=changes)
            messages.success(request, "Calibration entry saved.")
            return redirect("device_detail", pk=device.pk)
    else:
        form = CalibrationEventForm(instance=event)
    return render(request, "calibration/event_form.html", {"form": form, "device": device, "event": event})


@login_required
def event_delete(request, pk):
    event = get_object_or_404(CalibrationEvent, pk=pk)
    device = _get_device(request, event.device_id, edit=True)
    if request.method == "POST":
        delete_calibration_event(event, request.user)
        messages.success(request, "Calibration entry deleted.")
        return redirect("device_detail", pk=device.pk)
    return render(request, "calibration/confirm_delete.html", {
        "object": event, "what": "calibration entry",
        "warning": "Attached certificates are deleted too. The next calibration date is recalculated.",
        "cancel_url": reverse("device_detail", args=[device.pk]),
    })


# -------------------------------------------------------------- attachments


@login_required
@require_POST
def attachment_upload(request, pk):
    device = _get_device(request, pk, edit=True)
    form = AttachmentForm(request.POST, request.FILES)
    if form.is_valid():
        att = form.save(commit=False)
        att.device = device
        att.original_name = os.path.basename(att.file.name)
        att.uploaded_by = request.user
        att.save()
        audit(request.user, "file uploaded", att, device=device, changes={"File": (None, att.original_name)})
        messages.success(request, "File uploaded.")
    else:
        messages.error(request, "Upload failed: " + "; ".join(e for errs in form.errors.values() for e in errs))
    return redirect("device_detail", pk=device.pk)


@login_required
def attachment_download(request, pk):
    att = get_object_or_404(Attachment, pk=pk)
    _get_device(request, att.device_id)
    return _file_response(att.file, att.original_name)


@login_required
@require_POST
def attachment_delete(request, pk):
    att = get_object_or_404(Attachment, pk=pk)
    device = _get_device(request, att.device_id, edit=True)
    audit(request.user, "file deleted", att, device=device, changes={"File": (att.original_name, None)})
    att.file.delete(save=False)
    att.delete()
    messages.success(request, "File deleted.")
    return redirect("device_detail", pk=device.pk)


@login_required
def field_file_download(request, pk):
    value = get_object_or_404(DeviceFieldValue, pk=pk)
    _get_device(request, value.device_id)
    if not value.value_file:
        raise Http404
    return _file_response(value.value_file, value.display())


# ---------------------------------------------------------- brands & models


def _form_error_text(form):
    return "; ".join(e for errs in form.errors.values() for e in errs)


@login_required
@require_POST
def manufacturer_quick_add(request):
    """+ button on the device form. Returns the existing brand when the name is already taken."""
    name = " ".join(request.POST.get("name", "").split())
    existing = Manufacturer.objects.filter(name__iexact=name).first() if name else None
    if existing:
        return JsonResponse({"id": existing.pk, "name": existing.name, "existed": True})
    form = ManufacturerForm({"name": name})
    if not form.is_valid():
        return JsonResponse({"error": _form_error_text(form)}, status=400)
    brand = form.save(commit=False)
    brand.created_by = request.user
    brand.save()
    audit(request.user, "brand added", brand, changes={"Name": (None, brand.name)})
    return JsonResponse({"id": brand.pk, "name": brand.name, "existed": False})


@login_required
@require_POST
def model_quick_add(request):
    name = " ".join(request.POST.get("name", "").split())
    brand = Manufacturer.objects.filter(pk=request.POST.get("manufacturer") or None).first()
    if not brand:
        return JsonResponse({"error": "Select a manufacturer first."}, status=400)
    existing = DeviceModel.objects.filter(manufacturer=brand, name__iexact=name).first() if name else None
    if existing:
        return JsonResponse({"id": existing.pk, "name": existing.name, "manufacturer": brand.pk, "existed": True})
    form = DeviceModelForm({"manufacturer": brand.pk, "name": name})
    if not form.is_valid():
        return JsonResponse({"error": _form_error_text(form)}, status=400)
    model = form.save(commit=False)
    model.created_by = request.user
    model.save()
    audit(request.user, "model added", model, changes={"Manufacturer": (None, brand.name), "Name": (None, model.name)})
    return JsonResponse({"id": model.pk, "name": model.name, "manufacturer": brand.pk, "existed": False})


@editor_all_required
def brand_list(request):
    brands = Manufacturer.objects.annotate(device_count=Count("devices")).prefetch_related(
        Prefetch("models", queryset=DeviceModel.objects.annotate(device_count=Count("devices")))
    )
    return render(request, "calibration/brand_list.html", {
        "brands": brands,
        "brand_form": ManufacturerForm(prefix="brand"),
        "model_form": DeviceModelForm(prefix="model"),
    })


def _refresh_due_dates(devices):
    """Recompute next dates after an inherited interval changed."""
    for device in devices.filter(interval_override_months__isnull=True).select_related("manufacturer", "device_model"):
        device.refresh_calibration_dates()


@editor_all_required
def brand_detail(request, pk):
    brand = get_object_or_404(Manufacturer, pk=pk)
    return render(request, "calibration/lookup_detail.html", {
        "obj": brand, "what": "manufacturer",
        "fields": brand.fields.annotate(value_count=Count("values")),
        "device_count": brand.devices.count(),
        "models": brand.models.annotate(device_count=Count("devices")),
        "settings": SiteSettings.load(),
    })


@editor_all_required
def model_detail(request, pk):
    model = get_object_or_404(DeviceModel.objects.select_related("manufacturer"), pk=pk)
    return render(request, "calibration/lookup_detail.html", {
        "obj": model, "what": "model",
        "fields": model.fields.annotate(value_count=Count("values")),
        "brand_fields": model.manufacturer.fields.all(),
        "device_count": model.devices.count(),
        "settings": SiteSettings.load(),
    })


@editor_all_required
def brand_edit(request, pk=None):
    brand = get_object_or_404(Manufacturer, pk=pk) if pk else Manufacturer(created_by=request.user)
    form_class = ManufacturerEditForm if pk else ManufacturerForm
    before = snapshot(brand, form_class.Meta.fields) if pk else {}
    form = form_class(request.POST or None, instance=brand, prefix=None if pk else "brand")
    if request.method == "POST" and form.is_valid():
        changes = form_changes(form, before)
        brand = form.save()
        if "interval_months" in form.changed_data:
            _refresh_due_dates(brand.devices.all())
        audit(request.user, "brand edited" if pk else "brand added", brand, changes=changes)
        messages.success(request, f"Manufacturer “{brand.name}” saved.")
        return redirect("brand_detail", pk=brand.pk) if pk else redirect("brand_list")
    if request.method == "POST" and not pk:
        messages.error(request, _form_error_text(form))
        return redirect("brand_list")
    return render(request, "calibration/brand_form.html", {"form": form, "obj": brand, "what": "manufacturer"})


@editor_all_required
def model_edit(request, pk=None):
    model = get_object_or_404(DeviceModel, pk=pk) if pk else DeviceModel(created_by=request.user)
    form_class = DeviceModelEditForm if pk else DeviceModelForm
    before = snapshot(model, form_class.Meta.fields) if pk else {}
    form = form_class(request.POST or None, instance=model, prefix=None if pk else "model")
    if request.method == "POST" and form.is_valid():
        changes = form_changes(form, before)
        model = form.save()
        if pk and "manufacturer" in form.changed_data:
            # keep devices consistent: their brand follows the model
            model.devices.update(manufacturer=model.manufacturer)
        if {"manufacturer", "interval_months"} & set(form.changed_data):
            _refresh_due_dates(model.devices.all())
        audit(request.user, "model edited" if pk else "model added", model, changes=changes)
        messages.success(request, f"Model “{model.name}” saved.")
        return redirect("model_detail", pk=model.pk) if pk else redirect("brand_list")
    if request.method == "POST" and not pk:
        messages.error(request, _form_error_text(form))
        return redirect("brand_list")
    return render(request, "calibration/brand_form.html", {"form": form, "obj": model, "what": "model"})


@editor_all_required
@require_POST
def brand_delete(request, pk):
    brand = get_object_or_404(Manufacturer, pk=pk)
    if brand.devices.exists() or DeviceModel.objects.filter(manufacturer=brand, devices__isnull=False).exists():
        messages.error(request, f"“{brand.name}” is used by devices and cannot be deleted.")
        return redirect("brand_list")
    audit(request.user, "brand deleted", brand, changes={"Name": (brand.name, None)})
    brand.models.all().delete()
    brand.delete()
    messages.success(request, f"Manufacturer “{brand.name}” deleted.")
    return redirect("brand_list")


@editor_all_required
@require_POST
def model_delete(request, pk):
    model = get_object_or_404(DeviceModel, pk=pk)
    if model.devices.exists():
        messages.error(request, f"Model “{model.name}” is used by devices and cannot be deleted.")
        return redirect("brand_list")
    audit(request.user, "model deleted", model, changes={"Name": (model.name, None)})
    model.delete()
    messages.success(request, f"Model “{model.name}” deleted.")
    return redirect("brand_list")


# -------------------------------------------------- brand & model fields


def _owner_url(fd):
    if fd.device_model_id:
        return reverse("model_detail", args=[fd.device_model_id])
    return reverse("brand_detail", args=[fd.manufacturer_id])


@editor_all_required
def field_edit(request, pk=None, brand_pk=None, model_pk=None):
    if pk:
        fd = get_object_or_404(FieldDefinition.objects.select_related("manufacturer", "device_model"), pk=pk)
    elif model_pk:
        fd = FieldDefinition(device_model=get_object_or_404(DeviceModel, pk=model_pk))
    else:
        fd = FieldDefinition(manufacturer=get_object_or_404(Manufacturer, pk=brand_pk))
    owner = fd.owner
    if request.method == "POST":
        before = snapshot(fd, FieldDefinitionForm.Meta.fields) if pk else {}
        form = FieldDefinitionForm(request.POST, instance=fd)
        if form.is_valid():
            changes = form_changes(form, before)
            fd = form.save()
            audit(request.user, "field edited" if pk else "field added", fd, changes=changes)
            messages.success(request, f"Field '{fd.name}' saved.")
            return redirect(_owner_url(fd))
    else:
        initial = {} if pk else {"order": (owner.fields.count() + 1) * 10}
        form = FieldDefinitionForm(instance=fd, initial=initial)
    return render(request, "calibration/field_form.html", {
        "form": form, "owner": owner, "field": fd if pk else None, "cancel_url": _owner_url(fd),
    })


@editor_all_required
def field_delete(request, pk):
    fd = get_object_or_404(FieldDefinition.objects.select_related("manufacturer", "device_model"), pk=pk)
    count = fd.values.count()
    back = _owner_url(fd)
    if request.method == "POST":
        for v in fd.values.exclude(value_file=""):
            v.value_file.delete(save=False)
        audit(request.user, "field deleted", fd, changes={"Field": (fd.name, None)})
        fd.delete()
        messages.success(request, "Field deleted.")
        return redirect(back)
    return render(request, "calibration/confirm_delete.html", {
        "object": fd, "what": "field",
        "warning": f"{count} stored value(s) for this field will be deleted." if count else "",
        "cancel_url": back,
    })


# -------------------------------------------------------------------- users


@admin_required
def user_list(request):
    users = User.objects.annotate(device_count=Count("devices")).order_by("-is_active", "last_name", "username")
    return render(request, "calibration/user_list.html", {"users": users})


@admin_required
def user_create(request):
    if request.method == "POST":
        form = UserCreateForm(request.POST)
        if form.is_valid():
            user = form.save()
            audit(request.user, "user created", user, changes={"Role": (None, user.get_role_display())})
            messages.success(request, f"User {user.username} created.")
            return redirect("user_list")
    else:
        form = UserCreateForm()
    return render(request, "calibration/user_form.html", {"form": form, "obj": None})


@admin_required
def user_edit(request, pk):
    obj = get_object_or_404(User, pk=pk)
    if request.method == "POST":
        before = snapshot(obj, UserEditForm.Meta.fields)
        form = UserEditForm(request.POST, instance=obj)
        if form.is_valid():
            if obj == request.user and (
                not form.cleaned_data["is_active"]
                or (form.cleaned_data["role"] != User.ROLE_ADMIN and not obj.is_superuser)
            ):
                form.add_error(None, "You cannot remove your own administrator access or deactivate yourself.")
            else:
                changes = form_changes(form, before)
                form.save()
                if changes:
                    audit(request.user, "user edited", obj, changes=changes)
                messages.success(request, "User saved.")
                return redirect("user_list")
    else:
        form = UserEditForm(instance=obj)
    return render(request, "calibration/user_form.html", {"form": form, "obj": obj})


@admin_required
def user_password(request, pk):
    obj = get_object_or_404(User, pk=pk)
    form = SetPasswordForm(obj, request.POST or None)
    for f in form.fields.values():
        f.widget.attrs["class"] = "form-control"
    if request.method == "POST" and form.is_valid():
        form.save()
        audit(request.user, "password set", obj)
        messages.success(request, f"Password for {obj.username} changed.")
        return redirect("user_list")
    return render(request, "calibration/user_password.html", {"form": form, "obj": obj})


# ----------------------------------------------------------------- settings


@admin_required
def settings_view(request):
    settings = SiteSettings.load()
    if request.method == "POST" and "test_email" in request.POST:
        if not settings.smtp_configured:
            messages.error(request, "Save the SMTP host and From address first.")
        else:
            try:
                EmailMessage(
                    f"[{settings.site_name}] Test email",
                    "SMTP settings work. Calibration reminders can be delivered.",
                    settings.from_email,
                    [request.user.email],
                ).send()
                messages.success(request, f"Test email sent to {request.user.email}.")
            except Exception as e:
                messages.error(request, f"Sending failed: {e}")
        return redirect("settings")
    form = SiteSettingsForm(request.POST or None, instance=settings)
    if request.method == "POST" and form.is_valid():
        changes = form_changes(form, snapshot(SiteSettings.load(), SiteSettingsForm.Meta.fields))
        for key in ("SMTP password", "Remove stored SMTP password"):
            if key in changes:
                changes[key] = ("***", "***")
        form.save()
        if "default_interval_months" in form.changed_data:
            _refresh_due_dates(Device.objects.all())
        audit(request.user, "settings edited", settings, changes=changes)
        messages.success(request, "Settings saved.")
        return redirect("settings")
    return render(request, "calibration/settings.html", {
        "form": form, "settings": settings, "has_password": bool(settings.smtp_password),
    })


# ---------------------------------------------------------------- audit log


@editor_all_required
def audit_log(request):
    qs = AuditLog.objects.select_related("device")
    q = request.GET.get("q", "").strip()
    if q:
        qs = qs.filter(
            Q(username__icontains=q) | Q(object_repr__icontains=q) | Q(action__icontains=q)
            | Q(object_type__icontains=q)
        )
    return render(request, "calibration/audit_log.html", {"entries": qs[:500], "q": q})


# ------------------------------------------------------------------- import

IMPORT_SESSION_KEY = "device_import"


@editor_all_required
def import_upload(request):
    form = ImportUploadForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        try:
            headers, rows = importer.read_table(form.cleaned_data["file"])
        except Exception as e:
            form.add_error("file", f"Could not read the file: {e}")
        else:
            request.session[IMPORT_SESSION_KEY] = {"headers": headers, "rows": rows}
            return redirect("import_preview")
    return render(request, "calibration/import_upload.html", {
        "form": form, "targets": importer.TARGETS, "brands": Manufacturer.objects.prefetch_related("fields", "models__fields"),
    })


def _mapping_from_post(post, headers):
    mapping = {}
    for key in importer.TARGET_KEYS:
        raw = post.get(f"map_{key}", "")
        if raw.isdigit() and int(raw) < len(headers):
            mapping[key] = int(raw)
    return mapping


@editor_all_required
def import_preview(request):
    data = request.session.get(IMPORT_SESSION_KEY)
    if not data:
        messages.error(request, "Upload a file first.")
        return redirect("import_upload")
    headers, rows = data["headers"], data["rows"]
    if request.method == "POST":
        mapping = _mapping_from_post(request.POST, headers)
    else:
        mapping = importer.guess_mapping(headers)
    results = importer.validate(headers, rows, mapping)
    if request.method == "POST" and "apply" in request.POST:
        if any(r.action == "error" for r in results) and not request.POST.get("skip_errors"):
            messages.error(request, "Some rows have errors. Fix the file or tick 'skip rows with errors'.")
        else:
            created, updated = importer.apply(results, request.user)
            del request.session[IMPORT_SESSION_KEY]
            messages.success(request, f"Import finished: {created} created, {updated} updated.")
            return redirect("device_list")
    counts = {k: sum(1 for r in results if r.action == k) for k in ("create", "update", "error")}
    return render(request, "calibration/import_preview.html", {
        "headers": headers,
        "targets": [(k, label, req, mapping.get(k)) for k, label, req, _ in importer.TARGETS],
        "results": results,
        "counts": counts,
    })


# ---------------------------------------------------------- profile & ICS


@login_required
def profile(request):
    if request.method == "POST" and "regenerate" in request.POST:
        request.user.ics_token = new_ics_token()
        request.user.save(update_fields=["ics_token"])
        messages.success(request, "New calendar link created. The old link no longer works.")
        return redirect("profile")
    form = ProfileForm(request.POST or None, instance=request.user)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Profile saved.")
        return redirect("profile")
    ics_url = request.build_absolute_uri(reverse("ics_feed", args=[request.user.ics_token]))
    base = SiteSettings.load().base_url.rstrip("/")
    if base:
        ics_url = base + reverse("ics_feed", args=[request.user.ics_token])
    return render(request, "calibration/profile.html", {"form": form, "ics_url": ics_url})


def ics_feed(request, token):
    user = User.objects.filter(ics_token=token, is_active=True).first()
    if not user:
        raise Http404
    devices = visible_devices(user).select_related("owner", "manufacturer", "device_model")
    scope = request.GET.get("scope")
    if scope == "mine" or not user.can_view_all:
        devices = devices.filter(owner=user)
    settings = SiteSettings.load()
    body = build_calendar(devices, settings.base_url, name=f"{settings.site_name} calibrations")
    resp = HttpResponse(body, content_type="text/calendar; charset=utf-8")
    resp["Content-Disposition"] = 'inline; filename="calibrations.ics"'
    return resp
